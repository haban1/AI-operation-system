from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import select

from app.agent.service import AgentService
from app.api import knowledge
from app.api.deps import get_current_user
from app.api.schemas import ChatResponse
from app.core.config import Settings
from app.core.database import get_db
from app.knowledge.answers import validated_document_answer
from app.knowledge.parsers import parse_document
from app.knowledge.service import DocumentKnowledgeService
from app.llm.client import ChatCompletion
from app.llm.faq_orchestrator import FAQOrchestrator
from app.models.entities import FAQEntry, KnowledgeDocument, KnowledgeChunk, Role, User
from tests.llm_fakes import tool_call


class NoEmbedding:
    def embed(self, texts):
        return None


@pytest.fixture
def document_service():
    return DocumentKnowledgeService(embedding_client=NoEmbedding())


@pytest.fixture
def owner(db):
    user = User(username="document-admin", password_hash="x", role=Role.admin)
    db.add(user)
    db.commit()
    return user


def index(db, owner, service, text="Python零基础课程价格为1800元。报名截止日期为9月30日。", filename="课程介绍.txt"):
    doc = KnowledgeDocument(filename=filename, content=text.encode(), uploaded_by_id=owner.id)
    db.add(doc)
    db.commit()
    service.process(db, doc)
    return doc


def test_word_preserves_tables_and_paragraphs():
    from docx import Document
    doc = Document()
    doc.add_paragraph("Python入门课程")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "课程"
    table.cell(0, 1).text = "价格"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "1800元"
    output = BytesIO()
    doc.save(output)
    chunks = parse_document("课程.docx", output.getvalue())
    assert "1800元" in chunks[0]["text"]
    assert "价格" in chunks[0]["text"]
    assert "表 1" in chunks[0]["location"]


def test_empty_scanned_pdf_is_actionable_error():
    from pypdf import PdfWriter
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    output = BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="OCR"):
        parse_document("扫描.pdf", output.getvalue())


def test_pdf_page_locations_and_text():
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 50 700 Td (Python course costs 1800 yuan.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    chunks = parse_document("course.pdf", output.getvalue())
    assert chunks[0]["location"] == "第 1 页"
    assert "1800" in chunks[0]["text"]


def test_search_excludes_disabled_failed_and_legacy_faq(db, owner, document_service):
    doc = index(db, owner, document_service)
    db.add(FAQEntry(question="退款政策", answer="退款永久有效", course_code="general", intent_code="refund"))
    db.commit()
    hits = document_service.search(db, "Python价格")
    assert hits and hits[0]["filename"] == doc.filename
    assert hits[0]["material_id"] == doc.id
    assert document_service.search(db, "退款政策") == []
    doc.is_active = False
    db.commit()
    assert document_service.search(db, "Python价格") == []
    doc.is_active = True
    doc.status = "failed"
    db.commit()
    assert document_service.search(db, "Python价格") == []


def test_reindex_replaces_chunks_and_failed_reindex_not_searchable(db, owner, document_service):
    doc = index(db, owner, document_service)
    doc.content = "Python课程价格为2000元。".encode()
    db.commit()
    document_service.process(db, doc)
    chunks = list(db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == doc.id)))
    assert len(chunks) == 1 and "2000元" in chunks[0].text
    doc.content = b""
    db.commit()
    document_service.process(db, doc)
    assert doc.status == "failed"
    assert document_service.search(db, "Python价格") == []


def test_vector_retrieval_and_dimension_mismatch(db, owner):
    class Vectors:
        def embed(self, texts): return [[1.0, 0.0] for _ in texts]
    service = DocumentKnowledgeService(embedding_client=Vectors())
    service.settings = Settings(_env_file=None, dashscope_embedding_dimensions=2)
    doc = index(db, owner, service, "适合完全没有编程经验的同学。")
    assert doc.index_mode == "hybrid"
    assert service.search(db, "新手能参加吗")
    service.settings = Settings(_env_file=None, dashscope_embedding_dimensions=3)
    assert service.search(db, "Python报名") == []


def test_native_tool_chain_answers_from_document_with_citations(db, owner, document_service):
    doc = index(db, owner, document_service)
    calls = []
    class Client:
        def complete_with_tools(self, messages, **kwargs):
            calls.append(messages)
            if len(calls) == 1:
                return ChatCompletion(tool_calls=[tool_call("search_knowledge", {"query": "Python课程价格"})])
            import json
            tool_result = next(json.loads(m["content"]) for m in reversed(messages) if m["role"] == "tool")
            chunk = tool_result["document_chunks"][0]
            return ChatCompletion(content=f"Python课程价格为1800元。[来源:{chunk['chunk_id']}]")
    result = AgentService(document_service=document_service, faq_orchestrator=FAQOrchestrator(client=Client(), settings=Settings(_env_file=None))).answer(
        db, user=owner, message="Python课程多少钱？", session_id=None,
    )
    ChatResponse(**result)
    assert result["llm_used"] and result["tool_rounds"] == 1
    assert "1800元" in result["answer"] and "参考来源" in result["answer"]
    assert result["citations"][0]["filename"] == doc.filename
    assert result["matched_fact_ids"] == []
    assert result["matched_faq_id"] is None


def test_bad_citation_or_invented_amount_rejected():
    chunks = [{"chunk_id": "abc", "text": "价格1800元。"}]
    assert validated_document_answer("价格2000元。[来源:abc]", chunks, 4000) is None
    assert validated_document_answer("价格180元。[来源:abc]", chunks, 4000) is None
    assert validated_document_answer("价格1800元。[来源:wrong]", chunks, 4000) is None
    assert validated_document_answer("价格1800元。", chunks, 4000) is None
    assert validated_document_answer("价格1800元。[来源:abc]", chunks, 4000)


def test_document_api_lifecycle_and_permissions(db, owner, monkeypatch, document_service):
    monkeypatch.setattr(knowledge, "DocumentKnowledgeService", lambda: document_service)
    app = FastAPI()
    app.include_router(knowledge.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: owner
    client = TestClient(app)
    response = client.post("/knowledge/documents", files={"file": ("../课程.txt", "Python价格1800元".encode())})
    assert response.status_code == 201
    doc = response.json()
    assert doc["status"] == "ready" and doc["index_mode"] == "keyword"
    assert doc["filename"] == "课程.txt"
    path = f"/knowledge/documents/{doc['id']}"
    assert client.get(path).json()["chunks"][0]["text"] == "Python价格1800元"
    assert client.get(path + "/download").content == "Python价格1800元".encode()
    assert client.post("/knowledge/search", json={"query": "Python价格"}).json()["chunks"]
    assert client.patch(path, json={"is_active": False}).status_code == 200
    assert client.post("/knowledge/search", json={"query": "Python价格"}).json()["chunks"] == []
    assert client.post(path + "/reindex").json()["status"] == "ready"
    assert client.patch(path, json={"is_active": True}).status_code == 200
    student = User(username="student", password_hash="x", role=Role.student)
    app.dependency_overrides[get_current_user] = lambda: student
    for response in [client.get("/knowledge/documents"), client.get(path), client.delete(path), client.get(path + "/download"), client.post(path + "/reindex"), client.post("/knowledge/search", json={"query": "Python"})]:
        assert response.status_code == 403
    app.dependency_overrides[get_current_user] = lambda: owner
    assert client.delete(path).status_code == 204
    assert client.get(path).status_code == 404
    assert not list(db.scalars(select(KnowledgeChunk)))
    assert client.post("/knowledge/documents", files={"file": ("old.doc", b"data")}).status_code == 400
    assert client.post("/knowledge/documents", files={"file": ("empty.txt", b"")}).status_code == 400


def test_document_parse_failure_visible_in_api(db, owner, monkeypatch, document_service):
    monkeypatch.setattr(knowledge, "DocumentKnowledgeService", lambda: document_service)
    app = FastAPI()
    app.include_router(knowledge.router)
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: owner
    client = TestClient(app)
    response = client.post("/knowledge/documents", files={"file": ("broken.docx", b"not a zip")})
    assert response.status_code == 201 and response.json()["status"] == "failed"
    assert "损坏" in response.json()["error_message"]
