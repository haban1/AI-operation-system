from datetime import UTC, datetime
import re

from fastapi import APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status
from sqlalchemy import select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.api.deps import get_current_user
from app.api.schemas import (
    CourseDefinitionCreate,
    CourseDefinitionRead,
    CourseDefinitionUpdate,
    CourseIntentCreate,
    CourseIntentRead,
    CourseIntentUpdate,
    FAQCorrectionRequest,
    FAQCreate,
    FAQEmbeddingRefreshResult,
    FAQImportResult,
    FAQQueryLogRead,
    FAQRead,
    FAQTaxonomyRead,
    FAQUpdate,
)
from app.core.database import get_db
from app.core.permissions import require_role
from app.faq.embeddings import FAQEmbeddingService
from app.faq.facts import derive_fact_name
from app.faq.service import normalize_text
from app.intent.service import IntentService
from app.intent.taxonomy import (
    COURSES,
    DOMAINS,
    INTENTS,
    intent_id_for,
    intent_label,
    normalize_taxonomy,
)
from app.models.entities import (
    FAQEmbedding,
    FAQEntry,
    FAQQueryLog,
    CourseDefinition,
    IntentDefinition,
    Role,
    User,
)


router = APIRouter(prefix="/faqs", tags=["faqs"])
embedding_service = FAQEmbeddingService()
intent_service = IntentService()


@router.get("/taxonomy", response_model=FAQTaxonomyRead)
def get_faq_taxonomy(
    current_user: User = Depends(get_current_user),
) -> FAQTaxonomyRead:
    require_role(current_user, {Role.admin, Role.teacher})
    return FAQTaxonomyRead(
        domains=[{"code": item.code, "label": item.label} for item in DOMAINS],
        courses=[{"code": item.code, "label": item.label} for item in COURSES],
        intents=[{"code": item.code, "label": item.label} for item in INTENTS],
)


@router.get("/courses", response_model=list[CourseDefinitionRead])
def list_course_definitions(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[CourseDefinition]:
    require_role(current_user, {Role.admin, Role.teacher})
    intent_service.ensure_default_courses(db)
    return list(
        db.scalars(
            select(CourseDefinition)
            .where(CourseDefinition.is_deleted.is_(False))
            .order_by(CourseDefinition.created_at)
        )
    )


@router.post(
    "/courses",
    response_model=CourseDefinitionRead,
    status_code=status.HTTP_201_CREATED,
)
def create_course_definition(
    payload: CourseDefinitionCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CourseDefinition:
    require_role(current_user, {Role.admin, Role.teacher})
    course_code = payload.course_code.strip().lower()
    if not re.fullmatch(r"[a-z][a-z0-9_]{1,59}", course_code):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Course code must use lowercase letters, numbers and underscores",
        )
    display_name = payload.display_name.strip()
    if not display_name:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Course display name is required",
        )
    intent_service.ensure_default_courses(db)
    existing = db.get(CourseDefinition, course_code)
    if existing and not existing.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Course already exists",
        )
    aliases = _clean_list(payload.aliases)
    definition = existing or CourseDefinition(course_code=course_code)
    definition.display_name = display_name
    definition.description = (
        payload.description.strip()
        or f"{display_name}课程。"
    )
    definition.aliases = _clean_list([display_name, *aliases])
    definition.is_active = payload.is_active
    definition.is_deleted = False
    if existing:
        definition.version += 1
    else:
        db.add(definition)
    db.commit()
    db.refresh(definition)
    return definition


@router.patch("/courses/{course_code}", response_model=CourseDefinitionRead)
def update_course_definition(
    course_code: str,
    payload: CourseDefinitionUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> CourseDefinition:
    require_role(current_user, {Role.admin, Role.teacher})
    definition = db.get(CourseDefinition, course_code)
    if not definition or definition.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Course not found",
        )
    if payload.display_name is not None:
        display_name = payload.display_name.strip()
        if not display_name:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Course display name is required",
            )
        definition.display_name = display_name
    if payload.description is not None:
        definition.description = payload.description.strip()
    if payload.aliases is not None:
        definition.aliases = _clean_list(
            [definition.display_name, *payload.aliases]
        )
    if payload.is_active is not None:
        definition.is_active = payload.is_active
    definition.version += 1
    db.execute(
        update(IntentDefinition)
        .where(IntentDefinition.course_code == course_code)
        .values(course_name=definition.display_name)
    )
    db.execute(
        update(FAQEntry)
        .where(FAQEntry.course_code == course_code)
        .values(course_name=definition.display_name)
    )
    db.commit()
    db.refresh(definition)
    return definition


@router.delete(
    "/courses/{course_code}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_course_definition(
    course_code: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> None:
    require_role(current_user, {Role.admin, Role.teacher})
    definition = db.get(CourseDefinition, course_code)
    if not definition or definition.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Course not found",
        )
    definition.is_active = False
    definition.is_deleted = True
    definition.version += 1
    db.execute(
        update(IntentDefinition)
        .where(IntentDefinition.course_code == course_code)
        .values(is_active=False, is_deleted=True)
    )
    db.execute(
        update(FAQEntry)
        .where(FAQEntry.course_code == course_code)
        .values(is_active=False)
    )
    db.commit()


@router.get("/intents", response_model=list[CourseIntentRead])
def list_course_intents(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[IntentDefinition]:
    require_role(current_user, {Role.admin, Role.teacher})
    intent_service.ensure_default_definitions(db)
    return list(
        db.scalars(
            select(IntentDefinition)
            .where(
                IntentDefinition.is_deleted.is_(False),
            )
            .order_by(
                IntentDefinition.course_code,
                IntentDefinition.intent_id,
            )
        )
    )


@router.post(
    "/intents",
    response_model=CourseIntentRead,
    status_code=status.HTTP_201_CREATED,
)
def create_course_intent(
    payload: CourseIntentCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> IntentDefinition:
    require_role(current_user, {Role.admin, Role.teacher})
    course_code = payload.course_code.strip()
    intent_service.ensure_default_courses(db)
    course = db.get(CourseDefinition, course_code)
    if not course or course.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported course code",
        )
    intent_code = payload.intent_code.strip().lower()
    if not re.fullmatch(r"[a-z][a-z0-9_]{1,79}", intent_code):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Intent code must use lowercase letters, numbers and underscores",
        )
    display_name = payload.display_name.strip()
    if not display_name:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Intent display name is required",
        )
    positive_examples = _clean_list(payload.positive_examples)
    if not positive_examples:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="At least one positive example is required",
        )

    intent_service.ensure_default_definitions(db)
    intent_id = intent_id_for(course_code, intent_code)
    existing = db.get(IntentDefinition, intent_id)
    if existing and not existing.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Intent already exists",
        )
    definition = existing or IntentDefinition(
        intent_id=intent_id,
        domain="conversation" if course_code == "conversation" else "course",
        course_code=course_code,
        intent_code=intent_code,
    )
    definition.course_name = course.display_name
    definition.display_name = display_name
    definition.description = (
        payload.description.strip()
        or f"用户咨询{display_name}相关问题。"
    )
    definition.positive_examples = positive_examples
    definition.negative_examples = _clean_list(payload.negative_examples)
    definition.confusing_intents = []
    definition.risk_level = _risk_level(payload.risk_level)
    definition.is_active = payload.is_active
    definition.is_deleted = False
    if existing:
        definition.version += 1
    else:
        db.add(definition)
    db.commit()
    db.refresh(definition)
    return definition


@router.patch("/intents/{intent_id}", response_model=CourseIntentRead)
def update_course_intent(
    intent_id: str,
    payload: CourseIntentUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> IntentDefinition:
    require_role(current_user, {Role.admin, Role.teacher})
    definition = db.get(IntentDefinition, intent_id)
    if not definition or definition.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Intent not found",
        )
    if payload.display_name is not None:
        display_name = payload.display_name.strip()
        if not display_name:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Intent display name is required",
            )
        definition.display_name = display_name
    if payload.description is not None:
        definition.description = payload.description.strip()
    if payload.positive_examples is not None:
        positive_examples = _clean_list(payload.positive_examples)
        if not positive_examples:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="At least one positive example is required",
            )
        definition.positive_examples = positive_examples
    if payload.negative_examples is not None:
        definition.negative_examples = _clean_list(payload.negative_examples)
    if payload.risk_level is not None:
        definition.risk_level = _risk_level(payload.risk_level)
    if payload.is_active is not None:
        definition.is_active = payload.is_active
    definition.version += 1
    db.execute(
        update(FAQEntry)
        .where(FAQEntry.intent_id == intent_id)
        .values(category=definition.display_name)
    )
    db.commit()
    db.refresh(definition)
    return definition


@router.delete(
    "/intents/{intent_id}",
    status_code=status.HTTP_204_NO_CONTENT,
)
def delete_course_intent(
    intent_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> None:
    require_role(current_user, {Role.admin, Role.teacher})
    definition = db.get(IntentDefinition, intent_id)
    if not definition or definition.is_deleted:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Intent not found",
        )
    definition.is_active = False
    definition.is_deleted = True
    definition.version += 1
    db.execute(
        update(FAQEntry)
        .where(FAQEntry.intent_id == intent_id)
        .values(is_active=False)
    )
    db.commit()


@router.get("", response_model=list[FAQRead])
def list_faqs(
    include_inactive: bool = Query(True),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[FAQEntry]:
    require_role(current_user, {Role.admin, Role.teacher})
    stmt = select(FAQEntry).order_by(FAQEntry.priority.desc(), FAQEntry.updated_at.desc())
    stmt = stmt.where(
        FAQEntry.course_code.not_in(
            select(CourseDefinition.course_code).where(
                CourseDefinition.is_deleted.is_(True)
            )
        ),
        FAQEntry.intent_id.not_in(
            select(IntentDefinition.intent_id).where(
                IntentDefinition.is_deleted.is_(True)
            )
        ),
    )
    if not include_inactive:
        stmt = stmt.where(FAQEntry.is_active.is_(True))
    return list(db.scalars(stmt))


@router.get("/unmatched", response_model=list[FAQQueryLogRead])
def list_unmatched_questions(
    limit: int = Query(100, ge=1, le=500),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> list[FAQQueryLog]:
    require_role(current_user, {Role.admin, Role.teacher})
    return list(
        db.scalars(
            select(FAQQueryLog)
            .where(
                FAQQueryLog.review_status == "pending",
                FAQQueryLog.decision.in_(["CLARIFY", "OUT_OF_SCOPE"]),
            )
            .order_by(FAQQueryLog.created_at.desc())
            .limit(limit)
        )
    )


@router.patch("/query-logs/{log_id}/correction", response_model=FAQQueryLogRead)
def correct_query_log(
    log_id: str,
    payload: FAQCorrectionRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FAQQueryLog:
    require_role(current_user, {Role.admin, Role.teacher})
    log = db.get(FAQQueryLog, log_id)
    if not log:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Query log not found")

    corrected_entry = None
    embedding_changed = False
    if payload.corrected_faq_id:
        corrected_entry = db.get(FAQEntry, payload.corrected_faq_id)
        if not corrected_entry:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Corrected FAQ entry not found")
        log.corrected_faq_id = corrected_entry.id
        if payload.add_as_alias:
            normalized_alias = normalize_text(log.question)
            existing_values = [corrected_entry.question, *(corrected_entry.aliases or [])]
            if normalized_alias and all(normalize_text(value) != normalized_alias for value in existing_values):
                corrected_entry.aliases = [*(corrected_entry.aliases or []), log.question.strip()]
                embedding_changed = True

    log.correction_note = payload.correction_note.strip()
    log.review_status = "reviewed"
    log.reviewed_at = datetime.now(UTC)
    db.commit()
    db.refresh(log)
    if corrected_entry and embedding_changed:
        embedding_service.ensure_embeddings(db, [corrected_entry], force=True)
    return log


@router.post("/embeddings/rebuild", response_model=FAQEmbeddingRefreshResult)
def rebuild_faq_embeddings(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FAQEmbeddingRefreshResult:
    require_role(current_user, {Role.admin, Role.teacher})
    entries = list(db.scalars(select(FAQEntry).where(FAQEntry.is_active.is_(True))))
    vectors = embedding_service.ensure_embeddings(db, entries, force=True)
    return FAQEmbeddingRefreshResult(
        total=len(entries),
        available=len(vectors),
        model=embedding_service.model_name,
    )


@router.post("/import", response_model=FAQImportResult)
def import_faqs(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FAQImportResult:
    require_role(current_user, {Role.admin, Role.teacher})
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only XLSX files are supported")

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(file.file, read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Invalid XLSX file") from exc

    sheet = workbook["快捷回复列表"] if "快捷回复列表" in workbook.sheetnames else workbook.active
    existing_by_question = {entry.question: entry for entry in db.scalars(select(FAQEntry))}
    created = 0
    updated = 0
    skipped = 0
    changed_entries: list[FAQEntry] = []
    header_questions = {"问题", "标准问题", "用户问题", "关键词"}

    for values in sheet.iter_rows(values_only=True):
        cells = [str(value).strip() if value is not None else "" for value in values[:4]]
        category, question, aliases, answer = (cells + ["", "", "", ""])[:4]
        if not question or not answer or question in header_questions:
            skipped += 1
            continue
        alias_values = _clean_list(re_split_terms(aliases))
        course_code, course_name, intent_code = normalize_taxonomy(
            question=question,
            category=category,
            domain=category,
        )
        entry = existing_by_question.get(question)
        if entry:
            if not entry.fact_name.strip():
                entry.fact_name = derive_fact_name(question, category)
            entry.answer = answer
            entry.aliases = alias_values
            _set_taxonomy(
                entry,
                course_code=course_code,
                course_name=course_name,
                intent_code=intent_code,
                category=intent_label(intent_code),
                intent_id=intent_id_for(course_code, intent_code),
            )
            entry.is_active = True
            changed_entries.append(entry)
            updated += 1
            continue
        entry = FAQEntry(
            fact_name=derive_fact_name(question, category),
            question=question,
            answer=answer,
            aliases=alias_values,
            course_code=course_code,
            course_name=course_name,
            intent_code=intent_code,
            category=intent_label(intent_code),
            domain=(
                "conversation"
                if course_code == "conversation"
                else "course"
            ),
            intent_id=intent_id_for(course_code, intent_code),
            taxonomy_version="faq-taxonomy-v3",
            is_active=True,
            created_by_id=current_user.id,
        )
        db.add(entry)
        existing_by_question[question] = entry
        changed_entries.append(entry)
        created += 1

    db.commit()
    embedding_service.ensure_embeddings(db, changed_entries, force=True)
    workbook.close()
    return FAQImportResult(created=created, updated=updated, skipped=skipped)


@router.post("", response_model=FAQRead, status_code=status.HTTP_201_CREATED)
def create_faq(
    payload: FAQCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FAQEntry:
    require_role(current_user, {Role.admin, Role.teacher})
    fact_name = payload.fact_name.strip()
    question = (payload.question or fact_name).strip()
    answer = payload.answer.strip()
    image_url = _clean_image_url(payload.image_url)
    if not question or not answer:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Fact name and statement are required")
    (
        course_code,
        course_name,
        intent_code,
        category,
        intent_id,
    ) = _resolve_taxonomy(
        db,
        question=question,
        course_code=payload.course_code,
        course_name=payload.course_name,
        intent_code=payload.intent_code,
        category=payload.category,
        domain=payload.domain,
        intent_id=payload.intent_id,
    )
    entry = FAQEntry(
        fact_name=fact_name or derive_fact_name(question, category),
        question=question,
        answer=answer,
        image_url=image_url,
        aliases=_clean_list(payload.aliases),
        keywords=_clean_list(payload.keywords),
        course_code=course_code,
        course_name=course_name,
        intent_code=intent_code,
        category=category,
        domain=(
            "conversation"
            if course_code == "conversation"
            else "course"
        ),
        intent_id=intent_id,
        risk_level=_risk_level(payload.risk_level),
        handoff_required=payload.handoff_required,
        taxonomy_version="faq-taxonomy-v3",
        is_active=payload.is_active,
        priority=payload.priority,
        created_by_id=current_user.id,
    )
    db.add(entry)
    _commit_or_conflict(db)
    db.refresh(entry)
    embedding_service.ensure_embeddings(db, [entry], force=True)
    return entry


@router.patch("/{faq_id}", response_model=FAQRead)
def update_faq(
    faq_id: str,
    payload: FAQUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FAQEntry:
    require_role(current_user, {Role.admin, Role.teacher})
    entry = db.get(FAQEntry, faq_id)
    if not entry:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="FAQ entry not found")

    if payload.fact_name is not None:
        fact_name = payload.fact_name.strip()
        entry.fact_name = fact_name or derive_fact_name(
            payload.question if payload.question is not None else entry.question,
            entry.category,
        )
    if payload.question is not None:
        question = payload.question.strip()
        if not question:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Question is required")
        entry.question = question
    if payload.answer is not None:
        answer = payload.answer.strip()
        if not answer:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Answer is required")
        entry.answer = answer
    if payload.image_url is not None:
        entry.image_url = _clean_image_url(payload.image_url)
    if payload.aliases is not None:
        entry.aliases = _clean_list(payload.aliases)
    if payload.keywords is not None:
        entry.keywords = _clean_list(payload.keywords)
    taxonomy_changed = any(
        value is not None
        for value in (
            payload.course_code,
            payload.course_name,
            payload.intent_code,
            payload.category,
            payload.domain,
            payload.intent_id,
        )
    )
    if taxonomy_changed:
        (
            course_code,
            course_name,
            intent_code,
            category,
            intent_id,
        ) = _resolve_taxonomy(
            db,
            question=entry.question,
            course_code=payload.course_code if payload.course_code is not None else entry.course_code,
            course_name=payload.course_name if payload.course_name is not None else entry.course_name,
            intent_code=payload.intent_code if payload.intent_code is not None else entry.intent_code,
            category=payload.category if payload.category is not None else entry.category,
            domain=payload.domain if payload.domain is not None else entry.domain,
            intent_id=payload.intent_id if payload.intent_id is not None else entry.intent_id,
        )
        _set_taxonomy(
            entry,
            course_code=course_code,
            course_name=course_name,
            intent_code=intent_code,
            category=category,
            intent_id=intent_id,
        )
    if payload.risk_level is not None:
        entry.risk_level = _risk_level(payload.risk_level)
    if payload.handoff_required is not None:
        entry.handoff_required = payload.handoff_required
    if payload.is_active is not None:
        entry.is_active = payload.is_active
    if payload.priority is not None:
        entry.priority = payload.priority

    _commit_or_conflict(db)
    db.refresh(entry)
    embedding_service.ensure_embeddings(db, [entry], force=True)
    return entry


@router.post("/{faq_id}/image", response_model=FAQRead)
def upload_faq_image(
    faq_id: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> FAQEntry:
    require_role(current_user, {Role.admin, Role.teacher})
    entry = db.get(FAQEntry, faq_id)
    if not entry:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="FAQ entry not found",
        )

    content_type = (file.content_type or "").split(";", 1)[0].lower()
    if not content_type.startswith("image/"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only image files are supported",
        )
    image_data = file.file.read(10 * 1024 * 1024 + 1)
    if len(image_data) > 10 * 1024 * 1024:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Image file is larger than 10 MB",
        )
    if not image_data:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Image file is empty",
        )

    entry.image_data = image_data
    entry.image_url = ""
    entry.image_filename = re.sub(
        r"[\x00-\x1f]+",
        "",
        (file.filename or "customer-service-image").split("\\")[-1].split("/")[-1],
    )[:255]
    entry.image_content_type = content_type
    db.commit()
    db.refresh(entry)
    return entry


@router.delete("/{faq_id}/image", status_code=status.HTTP_204_NO_CONTENT)
def delete_faq_image(
    faq_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> None:
    require_role(current_user, {Role.admin, Role.teacher})
    entry = db.get(FAQEntry, faq_id)
    if not entry:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="FAQ entry not found",
        )
    entry.image_data = None
    entry.image_filename = ""
    entry.image_content_type = ""
    entry.image_url = ""
    db.commit()


@router.get("/{faq_id}/image")
def get_faq_image(faq_id: str, db: Session = Depends(get_db)) -> Response:
    entry = db.get(FAQEntry, faq_id)
    if not entry or not entry.image_data:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="FAQ image not found",
        )
    return Response(
        content=entry.image_data,
        media_type=entry.image_content_type or "application/octet-stream",
    )


@router.delete("/{faq_id}", status_code=status.HTTP_204_NO_CONTENT)
def disable_faq(
    faq_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> None:
    require_role(current_user, {Role.admin, Role.teacher})
    entry = db.get(FAQEntry, faq_id)
    if not entry:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="FAQ entry not found")
    entry.is_active = False
    db.commit()


@router.delete("/{faq_id}/permanent", status_code=status.HTTP_204_NO_CONTENT)
def permanently_delete_faq(
    faq_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
) -> None:
    require_role(current_user, {Role.admin})
    entry = db.get(FAQEntry, faq_id)
    if not entry:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="FAQ entry not found")

    db.execute(
        update(FAQQueryLog)
        .where(FAQQueryLog.matched_faq_id == faq_id)
        .values(matched_faq_id=None)
    )
    db.execute(
        update(FAQQueryLog)
        .where(FAQQueryLog.corrected_faq_id == faq_id)
        .values(corrected_faq_id=None)
    )
    embedding = db.get(FAQEmbedding, faq_id)
    if embedding:
        db.delete(embedding)
    db.delete(entry)
    db.commit()


def _clean_image_url(value: str | None) -> str:
    image_url = (value or "").strip()
    if not image_url:
        return ""
    if not re.fullmatch(r"https?://\S+", image_url, flags=re.IGNORECASE):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Image URL must be an HTTP or HTTPS URL",
        )
    return image_url


def _clean_list(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value.strip() for value in values if value.strip()))


def _set_taxonomy(
    entry: FAQEntry,
    *,
    course_code: str,
    course_name: str,
    intent_code: str,
    category: str,
    intent_id: str,
) -> None:
    entry.course_code = course_code
    entry.course_name = course_name
    entry.intent_code = intent_code
    entry.category = category
    entry.domain = (
        "conversation"
        if course_code == "conversation"
        else "course"
    )
    entry.intent_id = intent_id
    entry.taxonomy_version = "faq-taxonomy-v3"


def _resolve_taxonomy(
    db: Session,
    *,
    question: str,
    course_code: str = "",
    course_name: str = "",
    intent_code: str = "",
    category: str = "",
    domain: str = "",
    intent_id: str = "",
) -> tuple[str, str, str, str, str]:
    requested_course = course_code.strip()
    requested_intent = intent_code.strip()
    course = db.get(CourseDefinition, requested_course)
    if course and requested_intent:
        requested_id = intent_id_for(requested_course, requested_intent)
        definition = db.get(IntentDefinition, requested_id)
        if (
            definition
            and not definition.is_deleted
        ):
            return (
                definition.course_code,
                definition.course_name,
                definition.intent_code,
                definition.display_name,
                definition.intent_id,
            )

    normalized_course, normalized_name, normalized_intent = normalize_taxonomy(
        question=question,
        course_code=course_code,
        course_name=course_name,
        intent_code=intent_code,
        category=category,
        domain=domain,
        intent_id=intent_id,
    )
    normalized_id = intent_id_for(normalized_course, normalized_intent)
    definition = db.get(IntentDefinition, normalized_id)
    return (
        normalized_course,
        normalized_name,
        normalized_intent,
        definition.display_name if definition else intent_label(normalized_intent),
        normalized_id,
    )


def re_split_terms(value: str) -> list[str]:
    import re

    return re.split(r"[\n,，;；]+", value)


def _domain_value(value: str) -> str:
    normalized = normalize_text(value)
    return normalized or "general"


def _intent_value(value: str) -> str:
    return f"{_domain_value(value)}.faq"


def _risk_level(value: str) -> str:
    normalized = value.strip().lower()
    if normalized not in {"low", "medium", "high"}:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Risk level must be low, medium or high")
    return normalized


def _commit_or_conflict(db: Session) -> None:
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Question already exists") from exc
