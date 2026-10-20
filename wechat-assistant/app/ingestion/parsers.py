from pathlib import Path


def parse_material(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pptx":
        return parse_pptx(path)
    raise ValueError(f"Unsupported material format: {suffix}")


def parse_pdf(path: Path) -> list[dict[str, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append({"page_or_slide": f"page {index}", "text": text})
    return pages


def parse_docx(path: Path) -> list[dict[str, str]]:
    from docx import Document

    doc = Document(str(path))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip())
    return [{"page_or_slide": "document", "text": text}] if text.strip() else []


def parse_pptx(path: Path) -> list[dict[str, str]]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text)
        if texts:
            slides.append({"page_or_slide": f"slide {index}", "text": "\n".join(texts)})
    return slides


def chunk_sections(sections: list[dict[str, str]], *, max_chars: int = 1200, overlap: int = 120) -> list[dict[str, str]]:
    chunks: list[dict[str, str]] = []
    for section in sections:
        text = " ".join(section["text"].split())
        start = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            chunk_text = text[start:end].strip()
            if chunk_text:
                chunks.append({"page_or_slide": section["page_or_slide"], "text": chunk_text})
            if end == len(text):
                break
            start = max(0, end - overlap)
    return chunks

